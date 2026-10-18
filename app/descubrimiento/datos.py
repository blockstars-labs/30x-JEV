"""Lectura de datasets (CSV o Excel) y utilidades de texto, sin depender de cómo vengan las columnas."""

import csv
import io
import re
import unicodedata
from pathlib import Path

# Palabras vacías del español (y algunas del inglés): no distinguen a nadie.
VACIAS = set("""
a al algo algun alguna algunas alguno algunos ante antes aqui asi aun bien cada como con contra cual cuales
cuando de del desde donde dos el ella ellas ellos en entre era es esa esas ese eso esos esta estan estar
estas este esto estos fue ha hace hacer han hasta hay la las le les lo los mas me mi mis mucho muy nada
ni no nos nosotros o otra otras otro otros para pero poco por porque que quien quiero quisiera se sea ser si
sin sobre solo son su sus tambien tan tanto te tengo tenemos tiene tienen todo todos tu un una uno unos
vez y ya yo hola gracias saludos the and for with you your this that are was have from
""".split())

_PALABRA = re.compile(r"[a-zñ]{3,}")


def sin_tildes(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def normalizar(texto) -> str:
    """Minúsculas y sin tildes: 'Expandirnos a México' y 'expandirnos a mexico' cuentan igual."""
    return sin_tildes(str(texto or "")).lower().strip()


def terminos(texto) -> set[str]:
    """Palabras y pares de palabras significativos de un texto (para el proponente estadístico)."""
    palabras = [p for p in _PALABRA.findall(normalizar(texto)) if p not in VACIAS]
    pares = {f"{a} {b}" for a, b in zip(palabras, palabras[1:])}
    return set(palabras) | pares


def leer_dataset(ruta: str | Path, hoja: str | None = None) -> list[dict[str, str]]:
    """Lee un CSV (coma o punto y coma) o un Excel. Todo valor queda como texto; vacíos como ''."""
    ruta = Path(ruta)
    if ruta.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        hoja_ws = libro[hoja] if hoja else (libro["Perfiles"] if "Perfiles" in libro.sheetnames else libro.worksheets[0])
        filas = list(hoja_ws.iter_rows(values_only=True))
        if not filas:
            return []
        encabezado = [str(c).strip() if c is not None else f"columna_{i}" for i, c in enumerate(filas[0])]
        return [{encabezado[i]: _texto(v) for i, v in enumerate(f) if i < len(encabezado)}
                for f in filas[1:] if f and any(v not in (None, "") for v in f)]
    return leer_csv(ruta.read_text(encoding="utf-8-sig"))


def leer_csv(texto: str) -> list[dict[str, str]]:
    muestra = texto[:4096]
    delimitador = ";" if muestra.count(";") > muestra.count(",") else ","
    lector = csv.DictReader(io.StringIO(texto), delimiter=delimitador)
    return [{(k or "").strip(): (v or "").strip() for k, v in fila.items() if k is not None} for fila in lector]


def escribir_csv(ruta: str | Path, filas: list[dict], columnas: list[str] | None = None) -> None:
    columnas = columnas or list(filas[0].keys())
    with Path(ruta).open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=columnas)
        w.writeheader()
        for fila in filas:
            w.writerow({c: fila.get(c, "") for c in columnas})


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10]
    return str(v).strip()
